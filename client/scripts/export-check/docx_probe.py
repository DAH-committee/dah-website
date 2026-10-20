# python-docx 로 열어 보고 표, 병합, 쪽 설정을 JSON 으로 출력한다. 사용: python3 docx_probe.py 파일.docx
import sys, json
from docx import Document
from docx.oxml.ns import qn
def main(p):
    out = {'opened': False}
    try:
        d = Document(p)
    except Exception as e:
        out['error'] = f'{type(e).__name__}: {e}'
        print(json.dumps(out, ensure_ascii=False)); return
    out['opened'] = True
    out['paragraphs'] = len(d.paragraphs)
    out['sections'] = [{
        'wPt': round(s.page_width.pt, 1) if s.page_width else None,
        'hPt': round(s.page_height.pt, 1) if s.page_height else None,
        'orient': str(s.orientation),
        'marginPt': {k: round(getattr(s, k + '_margin').pt, 1) for k in ('top', 'bottom', 'left', 'right')},
    } for s in d.sections]
    tables = []
    for t in d.tables:
        tcs = t._tbl.findall('.//' + qn('w:tc'))
        gs = sum(1 for tc in tcs if tc.find(qn('w:tcPr')) is not None and tc.find(qn('w:tcPr')).find(qn('w:gridSpan')) is not None)
        vm = sum(1 for tc in tcs if tc.find(qn('w:tcPr')) is not None and tc.find(qn('w:tcPr')).find(qn('w:vMerge')) is not None)
        try:
            ncols = len(t.columns)
        except Exception:
            ncols = None
        tables.append({'rows': len(t.rows), 'cols': ncols, 'gridSpanCells': gs, 'vMergeCells': vm})
    out['tables'] = tables
    out['inlineShapes'] = len(d.inline_shapes)
    try:
        out['title'] = d.core_properties.title
    except Exception:
        pass
    print(json.dumps(out, ensure_ascii=False))
main(sys.argv[1])
